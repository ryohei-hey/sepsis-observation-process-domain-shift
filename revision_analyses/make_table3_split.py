"""
R2 (Reviewer 3, Major 4 + typesetting): split the 15-column domain-shift table into
  Table 3  Discrimination  (AUROC, AUPRC, their deltas, optimism gap)
  Table 4  Calibration     (Brier internal / external, CITL, calibration slope)
and DROP the ΔBrier column (prevalence-dependent; the null-model Brier differs by ~0.017
between cohorts, larger than most ΔBrier values).

Source of truth: outputs/outputs_for_manuscript/Table_domain_shift_performance.md
(regenerated 2026-07-29 by the full re-run; the docx is derived from the same rows).
Output: Table3_discrimination.docx / .md and Table4_calibration.docx / .md in the same folder,
written with the same python-docx layout as the notebook's save_table_as_docx().

R3 (Reviewer 3, Minor 3): Table 4 gains the internal calibration slope, read from
outputs/outputs_for_calibration_correction/STable_internal_external_slope_r3.md
(wp_o_internal_slope.py), so the S5 slope difference-in-differences can be reconstructed.
R3 (Language and presentation): CIs are written "(−0.051 to −0.027)" with a true minus sign
instead of "(-0.051–-0.027)". Formatting only; every number is asserted unchanged.
"""
import re
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
# The analysis notebook writes to output/ in the public repository and to outputs/ in the
# original project layout; Tables 3 and 4 are written next to their source.
OUT = next((d for d in (ROOT / "output" / "outputs_for_manuscript", ROOT / "outputs" / "outputs_for_manuscript")
            if (d / "Table_domain_shift_performance.md").exists()), ROOT / "output" / "outputs_for_manuscript")
SRC = OUT / "Table_domain_shift_performance.md"
INTERNAL_SLOPE = ROOT / "outputs" / "outputs_for_calibration_correction" / "STable_internal_external_slope_r3.md"
NUM = re.compile(r"[-−+]?\d+\.\d+")


def parse_md(path):
    lines = [l.rstrip("\n") for l in open(path, encoding="utf-8") if l.strip().startswith("|")]
    rows = [[c.strip() for c in l.strip("|").split("|")] for l in lines]
    return pd.DataFrame(rows[2:], columns=rows[0])


def numbers(s):
    return [float(x.replace("−", "-")) for x in NUM.findall(s)]


def format_ci(cell):
    """'-0.039 (-0.051–-0.027)' -> '−0.039 (−0.051 to −0.027)'; numbers unchanged."""
    out = re.sub(r"\(([^()]+?)–([^()]+?)\)", r"(\1 to \2)", cell)
    out = re.sub(r"(?<![\w.])-(?=\d)", "−", out)
    assert numbers(out) == numbers(cell), (cell, out)
    return out


def save_table_as_docx(df, footnotes_str, filepath):
    doc = Document()
    table = doc.add_table(rows=len(df) + 1, cols=len(df.columns), style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, col in enumerate(df.columns):
        cell = table.rows[0].cells[j]
        cell.text = str(col)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True; r.font.size = Pt(8)
    for i, (_, row) in enumerate(df.iterrows()):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(footnotes_str)
    r.font.size = Pt(8); r.italic = True
    doc.save(filepath)


def main():
    df = parse_md(SRC)
    cols = list(df.columns)
    print("source columns:", cols)
    # R2 (Reviewer 3, Language 3, second round): "Feature set" and "# Predictors" repeat
    # Table 1 / S1 Table and are the two widest columns, which starved the rest and made
    # Word break the remaining headings mid-word. Model ID keys the reader back to Table 1.
    key = ["Algorithm", "Model ID"]

    disc_cols = key + [
        "AUROC (Internal, 95% CI)", "AUROC (External, 95% CI)", "ΔAUROC (External − Internal, 95% CI)",
        "AUPRC (Internal, 95% CI)", "AUPRC (External, 95% CI)", "ΔAUPRC (External − Internal, 95% CI)",
        "Overfitting gap (Train − Internal AUROC)",
    ]
    cal_cols = key + [
        "Brier (Internal, 95% CI)", "Brier (External, 95% CI)",
        "Calibration-in-the-large in external (offset model, 95% CI)",
        "Calibration slope in external (95% CI)",
    ]
    missing = [c for c in disc_cols + cal_cols if c not in cols]
    assert not missing, missing

    disc = df[disc_cols].rename(columns={
        "ΔAUROC (External − Internal, 95% CI)": "ΔAUROC (95% CI)",
        "ΔAUPRC (External − Internal, 95% CI)": "ΔAUPRC (95% CI)",
        "Overfitting gap (Train − Internal AUROC)": "Optimism (Train − Internal AUROC)",
    })
    # The source md lost the trailing zeros of the optimism column. The notebook wrote it
    # with f"{overfitting_gap:.3f}" and its stored cell output still reads 0.050 and 0.070
    # (script/03_delivation_and_external_validation.ipynb), but the markdown round-trip turned
    # those strings into floats and back, so Table 3 printed 0.05 and 0.07 among otherwise
    # 3-decimal values. Restore the notebook's formatting; no value changes.
    opt_col = "Optimism (Train − Internal AUROC)"
    src_opt = df["Overfitting gap (Train − Internal AUROC)"].tolist()
    disc[opt_col] = [f"{float(v):.3f}" for v in src_opt]
    assert all(abs(float(a) - float(b)) < 1e-12 for a, b in zip(src_opt, disc[opt_col])),         "optimism reformatting changed a value"

    cal = df[cal_cols].rename(columns={
        "Calibration-in-the-large in external (offset model, 95% CI)": "Calibration-in-the-large, external (95% CI)",
        "Calibration slope in external (95% CI)": "Calibration slope, external (95% CI)",
    })
    # internal calibration slope (R3 Minor 3), placed before the external slope
    slp = parse_md(INTERNAL_SLOPE).set_index(["Algorithm", "Model ID"])
    keys = list(zip(cal["Algorithm"], cal["Model ID"]))
    assert all(slp.loc[k, "Table 4 external point"] == cal.loc[i, "Calibration slope, external (95% CI)"].split()[0]
               for i, k in zip(cal.index, keys)), "internal-slope file does not match Table 4 external slopes"
    cal.insert(cal.columns.get_loc("Calibration slope, external (95% CI)"), "Calibration slope, internal (95% CI)",
               [slp.loc[k, "Calibration slope, internal (95% CI)"] for k in keys])

    for d in (disc, cal):
        for c in d.columns[2:]:
            if c != opt_col:
                d[c] = [format_ci(v) for v in d[c]]

    disc_fn = (
        "Abbreviations: AUROC, area under the receiver operating characteristic curve; AUPRC, area under the "
        "precision–recall curve; CI, confidence interval. "
        "Internal validation: held-out MIMIC-IV test set (40%). External validation: independent eICU-CRD cohort. "
        "Δ = external − internal; negative Δ indicates performance degradation (domain shift). "
        "95% CIs: bootstrap resampling (B = 2,000, percentile method); Δ CIs from the bootstrap distribution of the difference. "
        "Outcome prevalence: internal = 16.3%; external = 13.9%. AUPRC depends mechanically on outcome prevalence, so part "
        "of each ΔAUPRC reflects the prevalence difference (S9 Table). "
        "Optimism = training AUROC − internal validation AUROC. "
        "Feature sets and predictor counts for each model specification are given in Table 1 and S1 Table. "
        "Preprocessing was identical for the internal and external datasets."
    )
    cal_fn = (
        "Abbreviations: CI, confidence interval. "
        "Internal validation: held-out MIMIC-IV test set (40%). External validation: independent eICU-CRD cohort. "
        "95% CIs: bootstrap resampling (B = 2,000, percentile method). "
        "Brier score: lower is better. Because the Brier score depends on outcome prevalence (a model predicting the "
        "base rate scores 0.136 internally and 0.120 externally, a difference of approximately 0.017), internal and external Brier scores are reported "
        "without their difference. "
        "Calibration-in-the-large estimated with an offset model, logit(Y) = α + logit(p̂), that is, with the calibration "
        "slope fixed at 1; the calibration slope estimated by unpenalized maximum likelihood from logit(Y) = α + β·logit(p̂), "
        "in the internal and in the external validation set. "
        "Ideal values: calibration-in-the-large = 0, slope = 1. "
        "Feature sets and predictor counts for each model specification are given in Table 1 and S1 Table."
    )

    save_table_as_docx(disc, disc_fn, OUT / "Table3_discrimination.docx")
    save_table_as_docx(cal, cal_fn, OUT / "Table4_calibration.docx")
    for name, d, fn in (("Table3_discrimination", disc, disc_fn), ("Table4_calibration", cal, cal_fn)):
        with open(OUT / f"{name}.md", "w", encoding="utf-8") as f:
            # disable_numparse: tabulate otherwise re-reads "0.050" as a float and drops the
            # trailing zero again, which is how the ragged optimism column arose.
            f.write(d.to_markdown(index=False, disable_numparse=True)); f.write("\n\n" + fn + "\n")
    print("rows", len(disc), len(cal), "-> Table3_discrimination / Table4_calibration (.docx/.md)")
    # null-model Brier check quoted in the footnote
    for p in (0.163, 0.139):
        print(f"  null Brier at prevalence {p}: {p*(1-p):.4f}")


if __name__ == "__main__":
    main()
